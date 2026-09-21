"""
data_pipeline/importers/investments/xlsx_probe.py
=================================================
Leitura mínima de um .xlsx para *decidir qual parser chamar*.

Existe um detector para os extratos da B3 e outro para os do Tesouro Direto.
Os dois precisam abrir o arquivo sem confiar nele, listar as abas e comparar
nomes sem acento — e duas cópias dessa base divergiriam na primeira correção
feita num lado só. As funções comuns moram aqui; cada sniffer guarda apenas as
assinaturas que são de fato suas.

Puro: recebe bytes, devolve dado. Não toca o banco e não importa Streamlit.
"""
from __future__ import annotations

import io
import unicodedata
from typing import Any


def norm(texto: str) -> str:
    """Caixa alta sem acento — 'Negociação' e 'Negociacao' colidem aqui."""
    sem_acento = unicodedata.normalize("NFKD", texto)
    sem_acento = "".join(c for c in sem_acento if not unicodedata.combining(c))
    return sem_acento.upper().strip()


def abrir(file_bytes: bytes):
    """Workbook somente-leitura, ou None se o arquivo não abrir.

    Quem chama é responsável por fechar. Arquivo corrompido, PDF renomeado ou
    openpyxl ausente caem todos no mesmo None: a resposta certa é recusar, e
    não entregar o arquivo a um parser no chute.
    """
    try:
        import openpyxl
    except ImportError:
        return None

    try:
        return openpyxl.load_workbook(
            io.BytesIO(file_bytes), read_only=True, data_only=True,
        )
    except Exception:  # noqa: BLE001 — arquivo corrompido ou não-xlsx
        return None


def fechar(workbook) -> None:
    """Fecha ignorando falha — o workbook já cumpriu o papel dele."""
    try:
        workbook.close()
    except Exception:  # noqa: BLE001
        pass


def sheet_names(file_bytes: bytes) -> list[str]:
    """Nomes das abas do .xlsx, ou lista vazia se o arquivo não abrir.

    Usado tanto pela detecção quanto pela mensagem de recusa — quando nada é
    reconhecido, a tela mostra o que o arquivo de fato tinha dentro.
    """
    wb = abrir(file_bytes)
    if wb is None:
        return []
    try:
        return list(wb.sheetnames)
    finally:
        fechar(wb)


def primeiras_linhas(sheet, limite: int = 12) -> list[tuple[Any, ...]]:
    """As `limite` primeiras linhas da aba, como tuplas de valores.

    Em `read_only=True` a planilha é um iterador; materializar só o topo é o
    suficiente para reconhecer um cabeçalho e evita varrer arquivo grande só
    para escolher o parser.
    """
    linhas: list[tuple[Any, ...]] = []
    for linha in sheet.iter_rows(values_only=True):
        linhas.append(linha)
        if len(linhas) >= limite:
            break
    return linhas
