"""Aba de aluguel do consolidado da XP: a Natureza decide o sinal.

Tickers, quantidades e valores são fictícios.
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest
from sqlalchemy import create_engine, event, text

from data_pipeline.importers.investments import xp_consolidado as xp

ARQUIVO = "relatorio-consolidado-mensal-2026-junho.xlsx"
DATA = date(2026, 6, 30)

CABECALHO_EMPRESTIMOS = [
    "Produto", "Instituição", "Natureza", "Número de Contrato", "Modalidade",
    "OPA", "Liquidação antecipada", "Taxa", "Comissão", "Data de registro",
    "Data de vencimento", "Quantidade", "Preço de Fechamento",
    "Valor Atualizado",
]


def _linha_emprestimo(produto, natureza, qtd, preco, valor):
    return [
        produto, "CORRETORA FICTICIA", natureza, "CONTRATO-TESTE",
        "Registro", "SIM", "SIM", 0.1, "-", "01/06/2026", "01/07/2026",
        qtd, preco, valor,
    ]


def _linha_acao(produto, qtd, preco, valor):
    linha = [None] * 14
    linha[0], linha[8], linha[12], linha[13] = produto, qtd, preco, valor
    return linha


def _planilha(acoes=(), emprestimos=()) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Posição - Ações"
    ws.append(["Produto"] + [f"col{i}" for i in range(1, 14)])
    for linha in acoes:
        ws.append(linha)
    ws = wb.create_sheet("Posição - Empréstimos")
    ws.append(CABECALHO_EMPRESTIMOS)
    for linha in emprestimos:
        ws.append(linha)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def engine(monkeypatch):
    eng = create_engine("sqlite://")

    # SAVEPOINT só funciona no pysqlite com o BEGIN emitido à mão.
    @event.listens_for(eng, "connect")
    def _conectar(dbapi_conn, _rec):
        dbapi_conn.isolation_level = None
        dbapi_conn.create_function("NOW", 0, lambda: "2026-06-30 00:00:00")

    @event.listens_for(eng, "begin")
    def _begin(conn):
        conn.exec_driver_sql("BEGIN")

    with eng.begin() as conn:
        conn.execute(text("""
            CREATE TABLE portfolio_position_snapshots (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id TEXT, portfolio_id TEXT, asset_id TEXT,
                report_date TEXT, quantity NUMERIC, market_price NUMERIC,
                market_value NUMERIC, invested_value NUMERIC,
                asset_name TEXT, asset_type TEXT, is_loaned BOOLEAN,
                institution TEXT, currency TEXT, country TEXT,
                source_system TEXT, source_table TEXT, source_id TEXT,
                imported_at TEXT,
                UNIQUE (portfolio_id, asset_id, report_date,
                        source_system, source_table, source_id)
            )
        """))

    monkeypatch.setattr(xp.settings, "OWNER_USER_ID", "user-1")
    monkeypatch.setattr(xp, "ensure_external_id_columns", lambda e: None)
    monkeypatch.setattr(xp, "_ensure_portfolio", lambda conn, uid: "portfolio-1")
    monkeypatch.setattr(
        xp, "get_or_create_asset",
        lambda conn, ticker, **kw: f"asset-{ticker}",
    )
    return eng


def _linhas(eng):
    with eng.connect() as conn:
        return [
            (r.asset_id, float(r.quantity), float(r.market_value),
             bool(r.is_loaned))
            for r in conn.execute(text("""
                SELECT asset_id, quantity, market_value, is_loaned
                  FROM portfolio_position_snapshots
                 ORDER BY asset_id, quantity
            """))
        ]


def _importar(eng, conteudo):
    return xp._parse_single((ARQUIVO, conteudo), eng)


def test_parser_assina_pela_natureza():
    wb = openpyxl.load_workbook(io.BytesIO(_planilha(emprestimos=[
        _linha_emprestimo("AAAA3 - FICTICIA A", "Doador", 40, 2.5, 100.0),
        _linha_emprestimo("BBBB3 - FICTICIA B", "Tomador", 10, 3.0, 30.0),
        _linha_emprestimo("CCCC3 - FICTICIA C", "Outra", 5, 1.0, 5.0),
    ])))
    doador, tomador, desconhecida = xp._parse_emprestimos(
        wb["Posição - Empréstimos"],
    )

    assert (doador["quantity"], doador["market_value"]) == (40, 100.0)
    assert (tomador["quantity"], tomador["market_value"]) == (-10, -30.0)
    assert tomador["market_price"] == 3.0
    assert doador["is_loaned"] and tomador["is_loaned"]
    assert desconhecida["ticker"] == "CCCC3"
    assert "Outra" in desconhecida["_erro"]


def test_tomador_grava_negativo_e_avisa(engine):
    resumo = _importar(engine, _planilha(emprestimos=[
        _linha_emprestimo("BBBB3 - FICTICIA B", "Tomador", 10, 3.0, 30.0),
    ]))

    assert _linhas(engine) == [("asset-BBBB3", -10.0, -30.0, True)]
    assert resumo["positions_imported"] == 1
    assert any("BBBB3" in n and "tomado" in n
               for n in resumo["files_skipped_notes"])


def test_tomador_nao_sobrescreve_acao_de_mesma_quantidade(engine):
    conteudo = _planilha(
        acoes=[_linha_acao("AAAA3 - FICTICIA A", 10, 3.0, 30.0)],
        emprestimos=[
            _linha_emprestimo("AAAA3 - FICTICIA A", "Tomador", 10, 3.0, 30.0),
        ],
    )
    _importar(engine, conteudo)
    _importar(engine, conteudo)

    assert _linhas(engine) == [
        ("asset-AAAA3", -10.0, -30.0, True),
        ("asset-AAAA3", 10.0, 30.0, False),
    ]


def test_reimportar_corrige_o_tomador_gravado_positivo(engine):
    chave_antiga = xp.make_external_id(
        "xpcsl-snap", [DATA.isoformat(), "BBBB3", "stock", f"{10:.6f}"],
    )
    with engine.begin() as conn:
        conn.execute(text("""
            INSERT INTO portfolio_position_snapshots
                (user_id, portfolio_id, asset_id, report_date, quantity,
                 market_value, asset_type, is_loaned, source_system,
                 source_table, source_id)
            VALUES ('user-1', 'portfolio-1', 'asset-BBBB3', :rd, 10, 30,
                    'stock', 1, 'app4', 'xp_consolidado', :sid)
        """), {"rd": DATA, "sid": chave_antiga})

    _importar(engine, _planilha(emprestimos=[
        _linha_emprestimo("BBBB3 - FICTICIA B", "Tomador", 10, 3.0, 30.0),
    ]))

    assert _linhas(engine) == [("asset-BBBB3", -10.0, -30.0, True)]


def test_doador_mantem_a_chave_de_sempre(engine):
    _importar(engine, _planilha(emprestimos=[
        _linha_emprestimo("AAAA3 - FICTICIA A", "Doador", 40, 2.5, 100.0),
    ]))

    chave = xp.make_external_id(
        "xpcsl-snap", [DATA.isoformat(), "AAAA3", "stock", f"{40:.6f}"],
    )
    with engine.connect() as conn:
        gravada = conn.execute(text(
            "SELECT source_id FROM portfolio_position_snapshots"
        )).scalar_one()
    assert gravada == chave
    assert _linhas(engine) == [("asset-AAAA3", 40.0, 100.0, True)]


def test_natureza_desconhecida_nao_grava_e_registra(engine):
    resumo = _importar(engine, _planilha(emprestimos=[
        _linha_emprestimo("CCCC3 - FICTICIA C", None, 5, 1.0, 5.0),
        _linha_emprestimo("AAAA3 - FICTICIA A", "Doador", 40, 2.5, 100.0),
    ]))

    assert _linhas(engine) == [("asset-AAAA3", 40.0, 100.0, True)]
    assert resumo["rows_skipped"] == 1
    assert any("CCCC3" in e and "(vazia)" in e for e in resumo["errors"])
